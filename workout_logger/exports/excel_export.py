from __future__ import annotations

from pathlib import Path
from collections import OrderedDict
from decimal import Decimal, InvalidOperation
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from ..database import Database
from ..notation import format_entry


def _duration_text(value):
    if value is None:
        return "Not recorded"
    try:
        duration = Decimal(str(value))
        if not duration.is_finite() or duration < 0:
            return "Not recorded"
        seconds = int(duration)
    except (InvalidOperation, TypeError, ValueError):
        return "Not recorded"
    hours, rest = divmod(seconds, 3600)
    minutes, seconds = divmod(rest, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


class _ExportSnapshot:
    def __init__(self, connection):
        self.connection = connection

    def fetchall(self, sql, params=()):
        return self.connection.execute(sql, params).fetchall()


class ExcelExporter:
    def __init__(self, db: Database):
        self.db = db

    def export(self, output_path: Path | str) -> Path:
        output = Path(output_path).expanduser()
        protected = {self.db.path.resolve(), Path(str(self.db.path) + "-wal").resolve(), Path(str(self.db.path) + "-shm").resolve()}
        if output.resolve() in protected:
            raise ValueError("Excel export must use a different path from the workout database")
        output.parent.mkdir(parents=True, exist_ok=True)

        # Every sheet must describe the same committed state, even if a second
        # app instance saves a workout while the workbook is being assembled.
        with self.db.connect() as connection:
            connection.execute("BEGIN")
            renderer = ExcelExporter(_ExportSnapshot(connection))
            wb = Workbook()
            ws = wb.active
            ws.title = "Workouts"
            renderer._workouts_sheet(ws)
            renderer._exercise_history_sheet(wb.create_sheet("Exercise History"))
            renderer._bodyweight_sheet(wb.create_sheet("Bodyweight"))
            renderer._raw_sheet(wb.create_sheet("Raw Data"))

        for sheet in wb.worksheets:
            # Workout metadata occupies five rows above the exercise results.
            sheet.freeze_panes = "B6" if sheet.title == "Workouts" else "A2"
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        # Names, notes and custom results are literal log text.
                        cell.data_type = "s"
            self._autosize(sheet)

        descriptor, name = tempfile.mkstemp(prefix=f".{output.name}.", suffix=".tmp", dir=output.parent)
        os.close(descriptor)
        tmp = Path(name)
        try:
            wb.save(tmp)
            tmp.replace(output)
        finally:
            wb.close()
            tmp.unlink(missing_ok=True)
        return output

    def _sessions(self):
        return [dict(r) for r in self.db.fetchall(
            """
            SELECT s.*, COALESCE(s.template_label_snapshot,t.label) AS template_label,
                   CASE WHEN t.active=0 THEN 'Removed Workout ' ELSE 'Workout ' END ||
                   COALESCE(s.template_label_snapshot,t.label,'—') AS workout_name
            FROM workout_sessions s
            LEFT JOIN workout_templates t ON t.id=s.template_id
            WHERE s.status IN ('complete','incomplete')
            ORDER BY s.session_date, s.id
            """
        )]

    def _workouts_sheet(self, ws):
        sessions = self._sessions()
        ws.cell(1, 1, "Exercise")
        ws.cell(2, 1, "Bodyweight")
        ws.cell(3, 1, "Status")
        ws.cell(4, 1, "Workout")
        ws.cell(5, 1, "Workout duration")
        for row in range(1, 6):
            ws.cell(row, 1).font = Font(bold=True)

        for col, session in enumerate(sessions, start=2):
            ws.cell(1, col, session["session_date"])
            ws.cell(2, col, f"{session['bodyweight']} kg")
            ws.cell(3, col, session["status"])
            ws.cell(4, col, session["workout_name"])
            ws.cell(5, col, _duration_text(session.get("workout_duration_seconds")))
            ws.cell(1, col).font = Font(bold=True)

        # Stable union of exercise snapshot names, then current display order for never-used exercises.
        names = OrderedDict()
        for r in self.db.fetchall(
            """
            SELECT we.exercise_name_snapshot
            FROM workout_entries we
            JOIN workout_sessions s ON s.id=we.session_id
            WHERE s.status IN ('complete','incomplete')
            ORDER BY s.session_date, s.id, we.exercise_order
            """
        ):
            names.setdefault(r["exercise_name_snapshot"], None)
        for r in self.db.fetchall("SELECT name FROM exercises ORDER BY display_order, id"):
            names.setdefault(r["name"], None)

        name_rows = {name: idx for idx, name in enumerate(names.keys(), start=6)}
        for name, row_idx in name_rows.items():
            ws.cell(row_idx, 1, name)

        session_col = {s["id"]: i for i, s in enumerate(sessions, start=2)}
        entries = self.db.fetchall(
            """
            SELECT we.* FROM workout_entries we
            JOIN workout_sessions s ON s.id=we.session_id
            WHERE s.status IN ('complete','incomplete')
            ORDER BY s.session_date, s.id, we.exercise_order
            """
        )
        for raw in entries:
            e = dict(raw)
            row_idx = name_rows[e["exercise_name_snapshot"]]
            col_idx = session_col[e["session_id"]]
            ws.cell(row_idx, col_idx, format_entry(e))

    def _exercise_history_sheet(self, ws):
        headers = ["Date", "Workout", "Exercise", "Performance", "Bodyweight", "Notes"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        rows = self.db.fetchall(
            """
            SELECT we.*, s.session_date, s.bodyweight AS session_bodyweight,
                   CASE WHEN t.active=0 THEN 'Removed Workout ' ELSE 'Workout ' END ||
                   COALESCE(s.template_label_snapshot,t.label,'—') AS workout_name
            FROM workout_entries we
            JOIN workout_sessions s ON s.id=we.session_id
            LEFT JOIN workout_templates t ON t.id=s.template_id
            WHERE s.status IN ('complete','incomplete')
            ORDER BY we.exercise_name_snapshot, s.session_date, s.id
            """
        )
        for raw in rows:
            e = dict(raw)
            ws.append([
                e["session_date"],
                e["workout_name"],
                e["exercise_name_snapshot"],
                format_entry(e),
                e["session_bodyweight"],
                e["notes"],
            ])

    def _bodyweight_sheet(self, ws):
        ws.append(["Date", "Bodyweight (kg)", "Workout Status"])
        for c in ws[1]:
            c.font = Font(bold=True)
        rows = self.db.fetchall(
            """
            SELECT bw.entry_date, bw.bodyweight, s.status
            FROM bodyweight_entries bw
            JOIN workout_sessions s ON s.id=bw.session_id
            WHERE s.status IN ('complete','incomplete')
            ORDER BY bw.entry_date, bw.id
            """
        )
        for r in rows:
            ws.append([r["entry_date"], r["bodyweight"], r["status"]])

    def _raw_sheet(self, ws):
        headers = [
            "session_id", "date", "workout", "status", "exercise_order", "exercise", "measurement_type",
            "unit", "bodyweight", "external_load", "assistance", "reps", "duration_seconds", "rounds",
            "rest_seconds", "resistance_value", "resistance_unit", "custom_result", "rir", "failure",
            "notes", "skipped", "workout_duration_seconds",
        ]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        rows = self.db.fetchall(
            """
            SELECT we.*, s.session_date, s.status, s.workout_duration_seconds,
                   CASE WHEN t.active=0 THEN 'Removed Workout ' ELSE 'Workout ' END ||
                   COALESCE(s.template_label_snapshot,t.label,'—') AS workout_name
            FROM workout_entries we
            JOIN workout_sessions s ON s.id=we.session_id
            LEFT JOIN workout_templates t ON t.id=s.template_id
            WHERE s.status IN ('complete','incomplete')
            ORDER BY s.session_date, s.id, we.exercise_order
            """
        )
        for r in rows:
            ws.append([
                r["session_id"], r["session_date"], r["workout_name"], r["status"], r["exercise_order"],
                r["exercise_name_snapshot"], r["measurement_type"], r["unit"], r["bodyweight"],
                r["external_load"], r["assistance"], r["reps"], r["duration_seconds"], r["rounds"],
                r["rest_seconds"], r["resistance_value"], r["resistance_unit"], r["custom_result"],
                r["rir_code"], r["failure"], r["notes"], r["skipped"], r["workout_duration_seconds"],
            ])

    def _autosize(self, ws):
        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                widths[cell.column] = min(max(widths.get(cell.column, 0), len(str(cell.value)) + 2), 60)
                cell.alignment = Alignment(vertical="top")
        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width
