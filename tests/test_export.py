import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from workout_logger.database import Database
from workout_logger.exports.excel_export import ExcelExporter
from workout_logger.models import WorkoutEntryData
from workout_logger.services.workout_service import WorkoutService


class ExportTests(unittest.TestCase):
    def test_removed_templates_keep_all_sessions_and_original_names_in_excel(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            db = Database(td / "workouts.db")
            service = WorkoutService(db)
            exercise_id = service.add_exercise("Curl", "external_weight")
            sessions = []
            for day in (1, 2):
                sid = service.start_new_template_session(f"2026-10-{day:02d}", "80")
                service.save_entry(sid, WorkoutEntryData(
                    exercise_id=exercise_id, exercise_name="Curl", measurement_type="external_weight",
                    external_load=Decimal("10"), reps=6, rir_code="0",
                ))
                service.finalize_session(sid)
                sessions.append(sid)
            service.remove_template(service.get_session(sessions[0])["template_id"])
            self.assertEqual([t["label"] for t in service.templates()], ["A"])
            output = ExcelExporter(db).export(td / "history.xlsx")
            wb = load_workbook(output)
            self.assertEqual(wb["Workouts"].cell(4, 2).value, "Removed Workout A")
            self.assertEqual(wb["Workouts"].cell(4, 3).value, "Workout B")
            self.assertEqual(wb["Exercise History"].cell(2, 2).value, "Removed Workout A")
            self.assertEqual(wb["Raw Data"].cell(2, 3).value, "Removed Workout A")
            self.assertEqual(wb["Exercise History"].max_row, 3)
            self.assertEqual(wb["Bodyweight"].max_row, 3)
            self.assertEqual(wb["Raw Data"].max_row, 3)
            wb.close()

    def test_end_to_end_database_and_excel(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            db = Database(td / "workouts.db")
            db.seed_exercises()
            service = WorkoutService(db)
            sid = service.start_new_template_session("2026-09-14", "84.6")
            exercises = {e["name"]: e for e in service.active_exercises()}

            ex = exercises["Super Wide-Grip Pull-Up"]
            service.save_entry(sid, WorkoutEntryData(
                exercise_id=ex["id"], exercise_name=ex["name"], measurement_type="bodyweight",
                unit="kg", bodyweight=Decimal("84.6"), reps=5, rir_code="0", notes="clean",
            ))
            ex = exercises["Chest Expander Above"]
            service.save_entry(sid, WorkoutEntryData(
                exercise_id=ex["id"], exercise_name=ex["name"], measurement_type="resistance_units",
                unit="springs", bodyweight=Decimal("84.6"), resistance_value=Decimal("4"),
                resistance_unit="springs", reps=7, rir_code="1",
            ))
            status = service.finalize_session(sid, force_incomplete=True)
            self.assertEqual(status, "incomplete")

            output = ExcelExporter(db).export(td / "history.xlsx")
            self.assertTrue(output.exists())
            wb = load_workbook(output)
            self.assertEqual(wb.sheetnames, ["Workouts", "Exercise History", "Bodyweight", "Raw Data"])
            ws = wb["Workouts"]
            self.assertEqual(ws.cell(1, 2).value, "2026-09-14")
            self.assertEqual(ws.cell(2, 2).value, "84.6 kg")
            values = [ws.cell(r, 2).value for r in range(5, ws.max_row + 1)]
            self.assertIn("BW × 5 @0 — clean", values)
            self.assertIn("4 springs × 7 @1", values)


if __name__ == "__main__":
    unittest.main()
