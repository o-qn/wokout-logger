import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from workout_logger.backups import BackupError
from workout_logger.database import Database
from workout_logger.exports.excel_export import ExcelExporter
from workout_logger.models import WorkoutEntryData
from workout_logger.services.workout_service import WorkoutService


class DurationExportTests(unittest.TestCase):
    def test_saved_duration_is_exported_and_automatic_backup_captures_saved_session(self):
        with tempfile.TemporaryDirectory() as td:
            service = WorkoutService(Database(Path(td) / 'history.db'))
            eid = service.add_exercise('Curl', 'external_weight')
            sid = service.start_new_template_session('2026-10-14', '80')
            start = datetime(2026, 10, 14, 15, tzinfo=timezone.utc)
            service.db.execute('UPDATE workout_sessions SET started_at=? WHERE id=?', (start.isoformat(), sid))
            service._timing_now = lambda: start + timedelta(seconds=90)
            service.save_entry(sid, WorkoutEntryData(exercise_id=eid, exercise_name='Curl', measurement_type='external_weight', external_load=Decimal('10'), reps=6, rir_code='0'))
            service.finalize_session(sid)
            self.assertTrue(any(b.reason == 'after-workout' and b.sessions == 1 and b.entries == 1 for b in service.backups.list_backups()))
            workbook = load_workbook(ExcelExporter(service.db).export(Path(td) / 'history.xlsx'))
            self.assertEqual(workbook['Workouts'].cell(5, 1).value, 'Workout duration')
            self.assertEqual(workbook['Workouts'].cell(5, 2).value, '00:01:30')
            raw = workbook['Raw Data']
            self.assertEqual(raw.cell(1, raw.max_column).value, 'workout_duration_seconds')
            self.assertEqual(float(raw.cell(2, raw.max_column).value), 90)
            workbook.close()

    def test_failed_after_workout_backup_does_not_lose_saved_workout(self):
        with tempfile.TemporaryDirectory() as td:
            service = WorkoutService(Database(Path(td) / 'history.db'))
            service.add_exercise('Curl', 'external_weight')
            sid = service.start_new_template_session('2026-10-14', '80')
            with patch.object(service.backups, 'create_backup', side_effect=BackupError('disk full')):
                self.assertEqual(service.finalize_session(sid, force_incomplete=True), 'incomplete')
            self.assertEqual(len(service.workout_history()), 1)
            self.assertEqual(service.backup_warning, 'disk full')


if __name__ == '__main__':
    unittest.main()
