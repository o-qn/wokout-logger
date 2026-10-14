import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from workout_logger.database import Database
from workout_logger.exports.excel_export import ExcelExporter, _duration_text
from workout_logger.models import WorkoutEntryData
from workout_logger.services.workout_service import WorkoutService


class ExportAuditTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.db = Database(self.folder / 'workouts.db')
        self.service = WorkoutService(self.db)
        self.eid = self.service.add_exercise('=1+1', 'custom')
        self.exporter = ExcelExporter(self.db)

    def tearDown(self):
        self.tmp.cleanup()

    def saved(self, day=1):
        sid = self.service.start_new_template_session(f'2026-10-{day:02d}', '80')
        self.service.save_entry(sid, WorkoutEntryData(exercise_id=self.eid, exercise_name='=1+1', measurement_type='custom', unit='units', custom_result='=2+2', notes='=3+3', rir_code='1'))
        self.service.finalize_session(sid)
        return sid

    def test_all_sheets_use_one_snapshot_during_concurrent_save(self):
        self.saved()
        original = ExcelExporter._sessions
        trigger = [False]
        def save_between_reads(renderer):
            rows = original(renderer)
            if not trigger[0]:
                trigger[0] = True
                self.saved(2)
            return rows
        output = self.folder / 'history.xlsx'
        with patch.object(ExcelExporter, '_sessions', save_between_reads):
            self.exporter.export(output)
        wb = load_workbook(output)
        self.assertEqual(len(self.service.workout_history()), 2)
        self.assertEqual(wb['Workouts'].max_column, 2)
        for name in ('Exercise History','Bodyweight','Raw Data'):
            self.assertEqual(wb[name].max_row, 2, name)
        wb.close()

    def test_formula_like_log_text_stays_literal_and_metadata_is_frozen(self):
        self.saved()
        output = self.exporter.export(self.folder / 'history.xlsx')
        wb = load_workbook(output, data_only=False)
        self.assertEqual(wb['Workouts'].freeze_panes, 'B6')
        for sheet, coordinate, expected in (
            ('Workouts','A6','=1+1'), ('Exercise History','C2','=1+1'),
            ('Exercise History','F2','=3+3'), ('Raw Data','R2','=2+2'), ('Raw Data','U2','=3+3'),
        ):
            cell = wb[sheet][coordinate]
            self.assertEqual(cell.value, expected)
            self.assertEqual(cell.data_type, 's')
        wb.close()

    def test_export_cannot_replace_database_or_journals(self):
        original = self.db.path.read_bytes()
        for name in (str(self.db.path),str(self.db.path)+'-wal',str(self.db.path)+'-shm'):
            with self.assertRaises(ValueError):
                self.exporter.export(name)
        self.assertEqual(self.db.path.read_bytes(), original)
        self.assertEqual(self.db.fetchone('PRAGMA integrity_check')[0], 'ok')

    def test_failed_save_keeps_previous_export_and_removes_partial_temp(self):
        self.saved()
        output = self.exporter.export(self.folder / 'history.xlsx')
        before = output.read_bytes()
        def fail_save(wb, path):
            Path(path).write_bytes(b'incomplete workbook')
            raise OSError('disk full')
        with patch.object(Workbook, 'save', fail_save):
            with self.assertRaises(OSError):
                self.exporter.export(output)
        self.assertEqual(output.read_bytes(), before)
        self.assertEqual(list(self.folder.glob('.history.xlsx.*.tmp')), [])

    def test_invalid_legacy_durations_do_not_break_export(self):
        for value in ('NaN','Infinity','not a duration','-1'):
            self.assertEqual(_duration_text(value), 'Not recorded')
        self.assertEqual(_duration_text('90.5'), '00:01:30')


if __name__ == '__main__':
    unittest.main()
